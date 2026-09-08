"""导出格式测试:九列 xlsx / Word 国内海外双格式 / 剪映草稿 zip."""
from __future__ import annotations

import io
import json
import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from agent.exporter import (
    export_storyboard_word,
    export_storyboard_xlsx,
)
from agent.jianying_exporter import (
    build_draft_content,
    export_jianying_draft_zip,
)
from agent.state import Asset
from agent.storyboard import Shot, Storyboard


def _storyboard() -> Storyboard:
    return Storyboard(
        title="测试剧本",
        emotion_type="复仇爽感",
        total_duration_seconds=30,
        shots=[
            Shot(shot_number=1, shot_type="特写", duration_seconds=15,
                 visual_description="0-3秒 特写女主A眼神\n3-6秒 中景办公室",
                 dialogue="我不会放弃", sound_effects="心跳声", bgm_mood="紧张",
                 camera_movement="推", fragment_id="S01-SC01-01", source_anchor="S01-P01",
                 characters=["char-01"], location_id="loc-01", props=["prop-01"]),
            Shot(shot_number=2, shot_type="中景", duration_seconds=15,
                 visual_description="0-3秒 调查桌前", dialogue="",
                 camera_movement="固定", fragment_id="S01-SC01-02"),
        ],
    )


def _assets() -> list[Asset]:
    return [
        Asset(id="char-01", kind="character", name="女主A", appearance="30岁白领"),
        Asset(id="loc-01", kind="location", name="办公室", appearance="冷白光"),
        Asset(id="prop-01", kind="prop", name="照片", appearance="旧木相框"),
    ]


def test_xlsx_nine_columns():
    from openpyxl import load_workbook
    buf = export_storyboard_xlsx("测试剧本", _storyboard(), assets=_assets())
    wb = load_workbook(buf)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    assert headers == ["镜号", "片段编号", "场景", "景别", "运镜", "时长(秒)",
                       "画面内容", "台词", "音效·BGM"], headers
    row1 = [c.value for c in ws[2]]
    assert row1[0] == 1
    assert "办公室" in row1[2]
    assert "我不会放弃" in row1[7]
    assert "紧张" in row1[8]
    print("[PASS] test_xlsx_nine_columns")


def test_word_domestic_and_overseas_formats():
    sb = _storyboard()
    domestic = export_storyboard_word("t", sb, assets=_assets(), fmt="国内")
    overseas = export_storyboard_word("t", sb, assets=_assets(), fmt="海外")

    # 粗验内容:docx 是 zip,直接比对 document.xml 里的标记
    import zipfile as zf
    def doc_text(buf: io.BytesIO) -> str:
        with zf.ZipFile(buf) as z:
            xml = z.read("word/document.xml").decode("utf-8")
        return xml

    d_text = doc_text(domestic)
    assert "△" in d_text          # 国内动作行前缀
    assert "♪" in d_text          # 国内音乐行前缀
    assert "国内" in d_text

    o_text = doc_text(overseas)
    assert "INT." in o_text or "EXT." in o_text   # 海外场景头
    assert "海外" in o_text
    assert "女主A" in o_text      # 资产名称出现在两版里
    print("[PASS] test_word_domestic_and_overseas_formats")


def test_jianying_draft_zip():
    sb = _storyboard()
    buf = export_jianying_draft_zip(sb, assets=_assets())
    with zipfile.ZipFile(buf) as zf:
        names = sorted(zf.namelist())
        assert names == ["README.txt", "draft_content.json", "draft_meta_info.json"]
        content = json.loads(zf.read("draft_content.json").decode("utf-8"))
        meta = json.loads(zf.read("draft_meta_info.json").decode("utf-8"))
        assert "剪映" in zf.read("README.txt").decode("utf-8")

    # 结构不变量
    video_tracks = [t for t in content["tracks"] if t["type"] == "video"]
    text_tracks = [t for t in content["tracks"] if t["type"] == "text"]
    assert len(video_tracks) == 1
    assert len(video_tracks[0]["segments"]) == 2
    # 视频轨时长 = 镜头时长之和(微秒)
    total_us = sum(s["target_timerange"]["duration"] for s in video_tracks[0]["segments"])
    assert total_us == 30 * 1_000_000
    assert content["duration"] == total_us
    # 视频素材名 = fragment_id(用户据此对号入座)
    mat_names = {m["material_name"] for m in content["materials"]["videos"]}
    assert mat_names == {"S01-SC01-01", "S01-SC01-02"}
    # 台词镜头产生 1 条字幕;无台词镜头不产生
    assert len(text_tracks[0]["segments"]) == 1
    assert meta["draft_name"] == "测试剧本"
    print("[PASS] test_jianying_draft_zip")


def test_jianying_subtitle_split_on_period():
    sb = Storyboard(title="t", emotion_type="e", total_duration_seconds=10,
                    shots=[Shot(shot_number=1, duration_seconds=10,
                                dialogue="第一句。第二句。第三句。",
                                fragment_id="S01-SC01-01")])
    content = build_draft_content(sb)
    text_track = next(t for t in content["tracks"] if t["type"] == "text")
    assert len(text_track["segments"]) == 3   # 按句拆 3 条字幕
    # 字幕时间衔接连续
    starts = [s["target_timerange"]["start"] for s in text_track["segments"]]
    durs = [s["target_timerange"]["duration"] for s in text_track["segments"]]
    assert starts == [0, durs[0], durs[0] + durs[1]]
    print("[PASS] test_jianying_subtitle_split_on_period")


if __name__ == "__main__":
    test_xlsx_nine_columns()
    test_word_domestic_and_overseas_formats()
    test_jianying_draft_zip()
    test_jianying_subtitle_split_on_period()
    print("\nAll exporter format tests passed.")
